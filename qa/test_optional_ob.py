"""Check optional fixed OB column without changing any real workbooks."""

from io import BytesIO
from copy import copy
from types import SimpleNamespace
from unittest.mock import patch
import unittest

import openpyxl

from app.services import excel_generator as generator


class OptionalOBTest(unittest.TestCase):
    def export(self, enabled=None, empty=False):
        host = SimpleNamespace(
            public_items=[] if empty else [{"item": "public", "gender": "All"}],
            self_paid_config={} if empty else {"A": [{"item": "paid", "gender": "All"}]},
            entry_reserve=SimpleNamespace(get=lambda: "100"),
            inst_var=SimpleNamespace(get=lambda: "H011-3"),
            app_settings={},
            get_effective_org_display=lambda item: (item["org_display"], False),
        )
        if enabled is not None:
            host.enable_ob_var = SimpleNamespace(get=lambda: enabled)
        items = {
            "public": dict(internal_code="003", sys_code="WBC", full_name="白血球", org_display="杏2"),
            "paid": dict(internal_code="107", sys_code="BUN", full_name="尿素氮", org_display="博2"),
        }
        stream = BytesIO()
        save = openpyxl.Workbook.save
        with (
            patch.object(generator, "_choose_output_path", return_value="synthetic.xlsx"),
            patch.object(generator, "find_db_item_robust_global", side_effect=items.get),
            patch.object(generator, "messagebox") as messages,
            patch.object(openpyxl.Workbook, "save", lambda book, path: save(book, stream)),
        ):
            generator.ExcelGenerationMixin.generate_excel(host)
        messages.showerror.assert_not_called()
        if not stream.getvalue():
            messages.showwarning.assert_called_once()
            messages.showinfo.assert_not_called()
            return None
        messages.showwarning.assert_not_called()
        messages.showinfo.assert_called_once()
        stream.seek(0)
        workbook = openpyxl.load_workbook(stream)
        self.addCleanup(workbook.close)
        self.addCleanup(stream.close)
        return workbook.active

    def test_disabled_and_legacy_host_preserve_existing_layout(self):
        for enabled in (False, None):
            with self.subTest(enabled=enabled):
                sheet = self.export(enabled)
                self.assertEqual([sheet.cell(3, col).value for col in range(13, 17)], ["105", "106", "107", "003"])
                self.assertEqual(sheet.max_column, 16)

    def test_enabled_ob_headers_and_all_formula_rows(self):
        sheet = self.export(True)
        self.assertEqual([sheet.cell(row, 15).value for row in range(1, 5)], ["杏1", None, "054", "OB"])
        self.assertEqual(sheet.max_row, 104)
        self.assertEqual(sheet.max_column, 17)
        for row in range(5, 105):
            self.assertEqual(sheet.cell(row, 15).value, f'=IF(AND($H{row}<>0,ISNUMBER(FIND("2",$J{row}))),"@","")')
        for row in range(1, 5):
            self.assertFalse(sheet.cell(row, 15).alignment.wrap_text)
            self.assertEqual(copy(sheet.cell(row, 15).alignment), copy(sheet.cell(row, 14).alignment))

    def test_only_ob_is_inserted_all_existing_cells_are_preserved(self):
        before, after = self.export(False), self.export(True)
        for row in before:
            for cell in row:
                moved = after.cell(cell.row, cell.column if cell.column < 15 else cell.column + 1)
                self.assertEqual(cell.value, moved.value, cell.coordinate)
                self.assertEqual(copy(cell.alignment), copy(moved.alignment), cell.coordinate)
                self.assertEqual(cell.number_format, moved.number_format, cell.coordinate)

    def test_ob_can_be_generated_without_other_scheduled_items(self):
        sheet = self.export(True, empty=True)
        self.assertEqual(sheet.max_column, 15)
        self.assertEqual(sheet["O4"].value, "OB")

    def test_empty_schedule_without_ob_still_warns(self):
        self.assertIsNone(self.export(False, empty=True))


if __name__ == "__main__":
    unittest.main()
