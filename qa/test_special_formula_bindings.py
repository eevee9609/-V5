"""Regression tests for exact item binding (110 must not match MAST GE110)."""

import json
import sys
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

import openpyxl

from app.services import excel_generator as generator


ITEMS = {
    "107": ("BUN", "尿素氮(BUN)"),
    "110": ("eGFR", "腎絲球過濾速率(eGFR)"),
    "199": ("MAST GE110", "過敏原檢查GE110"),
    "125": ("AMY", "澱粉酶"),
    "147": ("CA-125", "卵巢癌篩檢CA-125"),
    "101": ("A/G Ratio", "肝功能測試"),
    "200": ("MAST G101", "過敏原檢查G101"),
}
RULES = {"107": "腎", "110": "腎", "125": "胰", "101": "肝"}
SETTINGS = {
    "l_col_institutions": ["H011-3"],
    "special_formula_keywords_by_institution": {"H011-3": RULES},
}


def db_item(code):
    english, name = ITEMS[code]
    return dict(internal_code=code, sys_code=english, full_name=name, org_display="杏2")


def generate_formula_fixture(settings=None, inst_code="H011-3"):
    """Exercise the real export path using synthetic items and a temporary file."""
    host = SimpleNamespace(
        public_items=[],
        self_paid_config={
            "A": [{"item": code, "gender": "All"} for code in ("107", "110")],
            "V14": [{"item": "199", "gender": "All"}],
            "V7": [{"item": code, "gender": "All"} for code in ("125", "147", "101", "200")],
        },
        entry_reserve=SimpleNamespace(get=lambda: "1"),
        inst_var=SimpleNamespace(get=lambda: inst_code),
        app_settings=SETTINGS if settings is None else settings,
        get_effective_org_display=lambda item: (item["org_display"], False),
    )
    with tempfile.TemporaryDirectory(prefix="formula-binding-qa-") as directory:
        path = Path(directory) / "synthetic.xlsx"
        with (
            patch.object(generator, "_choose_output_path", return_value=str(path)),
            patch.object(generator, "find_db_item_robust_global", side_effect=db_item),
            patch.object(generator, "messagebox") as messages,
        ):
            generator.ExcelGenerationMixin.generate_excel(host)
            messages.showerror.assert_not_called()
            messages.showwarning.assert_not_called()
            messages.showinfo.assert_called_once()
        workbook = openpyxl.load_workbook(path)
        try:
            sheet = workbook["團檢大表"]
            return {
                sheet.cell(3, col).value: sheet.cell(5, col).value
                for col in range(15, sheet.max_column + 1)
            }
        finally:
            workbook.close()


class SpecialFormulaBindingsTest(unittest.TestCase):
    def test_generated_renal_formulas_only_bind_selected_items(self):
        formulas = generate_formula_fixture()
        for code in ("107", "110"):
            self.assertIn('FIND("腎",$K5)', formulas[code])
        self.assertNotIn('FIND("腎",$K5)', formulas["199"])

    def test_other_embedded_system_codes_do_not_leak(self):
        formulas = generate_formula_fixture()
        for selected, unrelated, keyword in (("125", "147", "胰"), ("101", "200", "肝")):
            self.assertIn(f'FIND("{keyword}",$K5)', formulas[selected])
            self.assertNotIn(f'FIND("{keyword}",$K5)', formulas[unrelated])

    def test_complete_legacy_english_and_name_keys_remain_supported(self):
        for key in ("eGFR", "腎絲球過濾速率(eGFR)", " 110 ", 110):
            with self.subTest(key=key):
                formulas = generate_formula_fixture({"special_formula_keywords": {key: "腎"}})
                self.assertIn('FIND("腎",$K5)', formulas["110"])
                self.assertNotIn('FIND("腎",$K5)', formulas["199"])

    def test_name_fragments_are_not_item_identities(self):
        formulas = generate_formula_fixture({"special_formula_keywords": {"GE110": "腎", "腎絲球": "炎"}})
        self.assertNotIn('FIND("腎",$K5)', formulas["199"])
        self.assertNotIn('FIND("炎",$K5)', formulas["110"])

    def test_institution_overrides_universal_without_leaking_to_other_institutions(self):
        settings = {
            "special_formula_keywords": {"110": "通用", "107": "共用"},
            "special_formula_keywords_by_institution": {" h011-3 ": {"110": "腎，炎"}},
        }
        formulas = generate_formula_fixture(settings)
        for word in ("腎", "炎"):
            self.assertIn(f'FIND("{word}",$K5)', formulas["110"])
        self.assertNotIn('FIND("通用",$K5)', formulas["110"])
        self.assertIn('FIND("共用",$K5)', formulas["107"])
        other = generate_formula_fixture(settings, "H999")
        self.assertIn('FIND("通用",$K5)', other["110"])
        self.assertNotIn('FIND("腎",$K5)', other["110"])

    def test_regular_package_conditions_unchanged(self):
        formulas = generate_formula_fixture()
        self.assertIn('FIND("A",UPPER(', formulas["110"])
        self.assertIn('FIND(",A,",', formulas["110"])
        self.assertIn('FIND("V14.",', formulas["199"])
        self.assertIn('FIND("V7.",', formulas["125"])
        self.assertIn('$H5<>"空號"', formulas["199"])
        self.assertIn('$K5<>"不抽血"', formulas["199"])

    def test_numeric_database_code_and_duplicate_keywords(self):
        item = dict(db_item("110"), internal_code=110)
        self.assertEqual(
            generator._get_item_special_keywords(item, {"110": ["腎"], "eGFR": ["腎", "炎"]}),
            ["腎", "炎"],
        )

    def test_no_binding_means_no_special_conditions(self):
        formulas = generate_formula_fixture({})
        for formula in formulas.values():
            for keyword in ("腎", "胰", "肝"):
                self.assertNotIn(f'FIND("{keyword}",$K5)', formula)


if __name__ == "__main__":
    if "--fixtures" in sys.argv:
        print(json.dumps(generate_formula_fixture(), ensure_ascii=False))
    else:
        unittest.main()
