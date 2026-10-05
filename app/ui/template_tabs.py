"""公費、自費套組與工作檔拆分頁面的介面控制。"""

import datetime
import json
import os
import re
import traceback
import tkinter as tk
from tkinter import filedialog, messagebox

import customtkinter as ctk
import openpyxl
import xlwings as xw
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.config import APP_ROOT, HEALTH_CHECK_ITEMS_FILE, UI_FONT
from app.data import (
    all_item_names,
    find_db_item_robust,
    get_sorted_items_list,
    load_health_check_database,
)
from app.services.column_routing import (
    column_routing_values,
    database_item_routing_values,
    override_org_display,
    parse_routing_rules,
    resolve_forced_destination,
)


CURRENT_DIR = str(APP_ROOT)
json_file = str(HEALTH_CHECK_ITEMS_FILE)
find_db_item_robust_global = find_db_item_robust


def load_json_database_directly(silent=True):
    return load_health_check_database(
        silent,
        on_warning=messagebox.showwarning,
        on_error=messagebox.showerror,
    )


class TemplateTabsMixin:

    def _selected_institution_data(self):
        """取得目前主畫面選定院所的分派規則。"""
        raw_value = self.inst_var.get().strip() if hasattr(self, "inst_var") else ""
        inst_code = raw_value.split(" - ")[0].strip().upper()
        return inst_code, self.inst_db.get(inst_code, {})

    def get_item_routing_override(self, db_item):
        """回傳目前院所對此資料庫項目的強制分派結果。"""
        _, inst_data = self._selected_institution_data()
        if not inst_data:
            return None, None, None
        return resolve_forced_destination(
            parse_routing_rules(inst_data.get("force_to_xinglian")),
            parse_routing_rules(inst_data.get("force_to_boren")),
            database_item_routing_values(db_item),
        )

    def get_effective_org_display(self, db_item):
        """取得匯出第一列應使用的歸屬標籤與是否被院所規則覆寫。"""
        destination, _, _ = self.get_item_routing_override(db_item)
        return override_org_display(db_item.get("org_display", ""), destination)

    def get_routed_item_display_label(self, display_label):
        """僅供介面顯示：保留原始資料，將開頭 [杏2] 改成 [博2]。"""
        db_item = find_db_item_robust_global(display_label)
        if not db_item:
            return display_label, False
        effective_org, overridden = self.get_effective_org_display(db_item)
        if not overridden:
            return display_label, False
        return re.sub(r"^\[[^\]]+\]", f"[{effective_org}]", display_label), True

    @staticmethod
    def _mark_routed_listbox_item(listbox, index, overridden):
        if overridden:
            listbox.itemconfig(index, foreground="#FFD54F")

    def build_public_tab(self, parent):
        left_outer = ctk.CTkFrame(parent, corner_radius=15)
        left_outer.pack(side="left", fill="both", expand=True, padx=10, pady=10)
        ctk.CTkLabel(left_outer, text=" 選擇要加入的公費項目 ", font=self.title_font).pack(pady=(10, 5))
        
        self.search_pub_var = tk.StringVar()
        self.entry_search_pub = ctk.CTkEntry(left_outer, textvariable=self.search_pub_var, font=self.default_font, placeholder_text="🔍 搜尋檢驗項目...", height=35)
        self.entry_search_pub.pack(fill="x", padx=15, pady=5)
        self.search_pub_var.trace_add("write", lambda *args: self.filter_items(self.entry_search_pub, self.listbox_pub, self.pub_sort_var))
        
        sort_pub_frame = ctk.CTkFrame(left_outer, fg_color="transparent")
        sort_pub_frame.pack(fill="x", padx=15, pady=(2, 5))
        ctk.CTkLabel(sort_pub_frame, text="排序方式:", font=(UI_FONT, 11, "bold")).pack(side="left")
        self.pub_sort_var = ctk.StringVar(value="預設 (原始順序)")
        self.pub_sort_combo = ctk.CTkOptionMenu(
            sort_pub_frame, values=["預設 (原始順序)", "系統代碼", "中/英文名稱", "英文簡寫", "歸屬單位", "檢驗類別"],
            variable=self.pub_sort_var, command=self.on_pub_sort_changed, height=26, font=(UI_FONT, 11)
        )
        self.pub_sort_combo.pack(side="left", padx=(5, 0), fill="x", expand=True)

        list_frame = ctk.CTkFrame(left_outer, fg_color="transparent")
        list_frame.pack(fill="both", expand=True, padx=15, pady=5)
        scroll = ctk.CTkScrollbar(list_frame)
        scroll.pack(side="right", fill="y")
        self.listbox_pub = tk.Listbox(list_frame, font=(UI_FONT, 12), selectmode="extended", yscrollcommand=scroll.set, 
                                      exportselection=False, bg="#2b2b2b", fg="white", relief="flat", highlightthickness=0, selectbackground="#1f538d", height=6)
        self.listbox_pub.pack(side="left", fill="both", expand=True)
        scroll.configure(command=self.listbox_pub.yview)
        
        gender_frame_pub = ctk.CTkFrame(left_outer, fg_color="transparent")
        gender_frame_pub.pack(anchor="center", pady=10)
        self.pub_gender_var = ctk.StringVar(value="All")
        ctk.CTkRadioButton(gender_frame_pub, text="不限男女", variable=self.pub_gender_var, value="All", font=self.default_font).pack(side="left", padx=10)
        ctk.CTkRadioButton(gender_frame_pub, text="👨 限定男", variable=self.pub_gender_var, value="Male", font=self.default_font).pack(side="left", padx=10)
        ctk.CTkRadioButton(gender_frame_pub, text="👩 限定女", variable=self.pub_gender_var, value="Female", font=self.default_font).pack(side="left", padx=10)
        
        self.listbox_pub.bind("<Double-Button-1>", lambda e: self.add_to_public())
        self._bind_picker_keyboard(self.listbox_pub, self.add_to_public, self.entry_search_pub)
        ctk.CTkButton(left_outer, text="➕ 加入至公費清單 (或雙擊項目)", font=self.default_font, height=40, command=self.add_to_public).pack(fill="x", padx=15, pady=(0, 15))
        
        right_outer = ctk.CTkFrame(parent, corner_radius=15)
        right_outer.pack(side="right", fill="both", expand=True, padx=10, pady=10)
        ctk.CTkLabel(right_outer, text=" 已排定的公費清單 ", font=self.title_font).pack(pady=(10, 5))
        
        spacer_top = ctk.CTkFrame(right_outer, height=75, fg_color="transparent")
        spacer_top.pack(fill="x")
        spacer_top.pack_propagate(False)
        
        workspace_frame = ctk.CTkFrame(right_outer, fg_color="transparent")
        workspace_frame.pack(fill="both", expand=True, padx=15, pady=5)
        list_frame_right = ctk.CTkFrame(workspace_frame, fg_color="transparent")
        list_frame_right.pack(side="left", fill="both", expand=True)
        scroll_right = ctk.CTkScrollbar(list_frame_right) 
        scroll_right.pack(side="right", fill="y")
        self.listbox_preview_pub = tk.Listbox(list_frame_right, font=(UI_FONT, 12), selectmode="extended", yscrollcommand=scroll_right.set, 
                                              bg="#1c2b36", fg="white", relief="flat", highlightthickness=0, selectbackground="#1f538d", height=6)
        self.listbox_preview_pub.pack(side="left", fill="both", expand=True)
        scroll_right.configure(command=self.listbox_preview_pub.yview)

        self.listbox_preview_pub.bind("<Double-Button-1>", lambda e: self.remove_from_public())
        self.listbox_preview_pub.bind("<Control-Up>", lambda e: (self.move_public_items(-1), "break")[1])
        self.listbox_preview_pub.bind("<Control-Down>", lambda e: (self.move_public_items(1), "break")[1])

        move_frame = ctk.CTkFrame(workspace_frame, width=38, fg_color="transparent", corner_radius=4)
        move_frame.pack(side="right", fill="y", padx=(8, 0))
        move_frame.pack_propagate(False)
        ctk.CTkButton(
            move_frame, text="↑", width=30, height=28, font=(UI_FONT, 16, "bold"),
            fg_color="#4A5568", hover_color="#2D3748", command=lambda: self.move_public_items(-1),
        ).pack(padx=2, pady=(4, 2))
        ctk.CTkButton(
            move_frame, text="↓", width=30, height=28, font=(UI_FONT, 16, "bold"),
            fg_color="#4A5568", hover_color="#2D3748", command=lambda: self.move_public_items(1),
        ).pack(padx=2, pady=2)

        btn_frame = ctk.CTkFrame(right_outer, fg_color="transparent")
        btn_frame.pack(fill="x", padx=15, pady=(5, 5))
        ctk.CTkButton(btn_frame, text="❌ 刪除選取", fg_color="#D9534F", hover_color="#C9302C", height=40, font=self.default_font, command=self.remove_from_public).pack(side="left", expand=True, padx=(0, 5))
        ctk.CTkButton(btn_frame, text="🔄 全部清空", fg_color="#5BC0DE", hover_color="#31B0D5", height=40, font=self.default_font, command=self.clear_public).pack(side="right", expand=True, padx=(5, 0))

        status_frame_pub = ctk.CTkFrame(right_outer, fg_color="transparent")
        status_frame_pub.pack(fill="x", padx=15, pady=(0, 10))
        self.lbl_count_pub = ctk.CTkLabel(status_frame_pub, text="當前已排定：0 項", font=(UI_FONT, 13, "bold"), text_color="#5BC0DE")
        self.lbl_count_pub.pack(side="right")
        
        self.refresh_listbox(self.listbox_pub, all_item_names, self.pub_sort_var.get())

    def build_self_paid_tab(self, parent):
        left_outer = ctk.CTkFrame(parent, corner_radius=15)
        left_outer.pack(side="left", fill="both", expand=True, padx=10, pady=10)
        ctk.CTkLabel(left_outer, text=" 選擇要加入的自費項目 ", font=self.title_font).pack(pady=(10, 5))
        
        pkg_frame = ctk.CTkFrame(left_outer, fg_color="transparent")
        pkg_frame.pack(fill="x", padx=15, pady=5)
        
        ctk.CTkLabel(pkg_frame, text="套組名稱:", font=self.default_font).pack(side="left", padx=(0, 5))
        self.entry_pkg = ctk.CTkEntry(pkg_frame, font=self.default_font, width=60, height=35)
        self.entry_pkg.pack(side="left")
        self.entry_pkg.insert(0, "A")
        
        # ⭐ 新增：單獨的 L 欄辨識輸入框
        ctk.CTkLabel(pkg_frame, text="單項加選代碼(選填):", font=self.default_font).pack(side="left", padx=(10, 5))
        self.entry_l_pkg = ctk.CTkEntry(pkg_frame, font=self.default_font, width=100, height=35, placeholder_text="如: B1, B2")
        self.entry_l_pkg.pack(side="left")
        
        self.search_self_var = tk.StringVar()
        self.entry_search_self = ctk.CTkEntry(left_outer, textvariable=self.search_self_var, font=self.default_font, placeholder_text="🔍 搜尋檢驗項目...", height=35)
        self.entry_search_self.pack(fill="x", padx=15, pady=(5, 5))
        self.search_self_var.trace_add("write", lambda *args: self.filter_items(self.entry_search_self, self.listbox_self, self.self_sort_var))
        
        sort_self_frame = ctk.CTkFrame(left_outer, fg_color="transparent")
        sort_self_frame.pack(fill="x", padx=15, pady=(2, 5))
        ctk.CTkLabel(sort_self_frame, text="排序方式:", font=(UI_FONT, 11, "bold")).pack(side="left")
        self.self_sort_var = ctk.StringVar(value="預設 (原始順序)")
        self.self_sort_combo = ctk.CTkOptionMenu(
            sort_self_frame, values=["預設 (原始順序)", "系統代碼", "中/英文名稱", "英文簡寫", "歸屬單位", "檢驗類別"],
            variable=self.self_sort_var, command=self.on_self_sort_changed, height=26, font=(UI_FONT, 11)
        )
        self.self_sort_combo.pack(side="left", padx=(5, 0), fill="x", expand=True)

        list_frame = ctk.CTkFrame(left_outer, fg_color="transparent")
        list_frame.pack(fill="both", expand=True, padx=15, pady=5)
        scroll = ctk.CTkScrollbar(list_frame)
        scroll.pack(side="right", fill="y")
        self.listbox_self = tk.Listbox(list_frame, font=(UI_FONT, 12), selectmode="extended", yscrollcommand=scroll.set, 
                                       exportselection=False, bg="#2b2b2b", fg="white", relief="flat", highlightthickness=0, selectbackground="#1f538d", height=6)
        self.listbox_self.pack(side="left", fill="both", expand=True)
        scroll.configure(command=self.listbox_self.yview)

        gender_frame_self = ctk.CTkFrame(left_outer, fg_color="transparent")
        gender_frame_self.pack(anchor="center", pady=10)
        self.self_paid_gender_var = ctk.StringVar(value="All")
        ctk.CTkRadioButton(gender_frame_self, text="不限男女", variable=self.self_paid_gender_var, value="All", font=self.default_font).pack(side="left", padx=10)
        ctk.CTkRadioButton(gender_frame_self, text="👨 限定男", variable=self.self_paid_gender_var, value="Male", font=self.default_font).pack(side="left", padx=10)
        ctk.CTkRadioButton(gender_frame_self, text="👩 限定女", variable=self.self_paid_gender_var, value="Female", font=self.default_font).pack(side="left", padx=10)

        self.listbox_self.bind("<Double-Button-1>", lambda e: self.add_to_self_paid())
        self._bind_picker_keyboard(self.listbox_self, self.add_to_self_paid, self.entry_search_self)
        ctk.CTkButton(left_outer, text="➕ 加入至自費套組 (或雙擊項目)", font=self.default_font, height=40, command=self.add_to_self_paid).pack(fill="x", padx=15, pady=(0, 15))
        
        right_outer = ctk.CTkFrame(parent, corner_radius=15)
        right_outer.pack(side="right", fill="both", expand=True, padx=10, pady=10)
        ctk.CTkLabel(right_outer, text=" 已排定的自費套組 ", font=self.title_font).pack(pady=(10, 5))
        
        spacer_top = ctk.CTkFrame(right_outer, height=120, fg_color="transparent")
        spacer_top.pack(fill="x")
        spacer_top.pack_propagate(False)

        workspace_frame = ctk.CTkFrame(right_outer, fg_color="transparent")
        workspace_frame.pack(fill="both", expand=True, padx=15, pady=5)
        list_frame_right = ctk.CTkFrame(workspace_frame, fg_color="transparent")
        list_frame_right.pack(side="left", fill="both", expand=True)
        scroll_right = ctk.CTkScrollbar(list_frame_right) 
        scroll_right.pack(side="right", fill="y")
        self.listbox_preview_self = tk.Listbox(list_frame_right, font=(UI_FONT, 12), selectmode="extended", yscrollcommand=scroll_right.set, 
                                               bg="#302717", fg="white", relief="flat", highlightthickness=0, selectbackground="#1f538d", height=6)
        self.listbox_preview_self.pack(side="left", fill="both", expand=True)
        scroll_right.configure(command=self.listbox_preview_self.yview)

        self.listbox_preview_self.bind("<Double-Button-1>", lambda e: self.remove_from_self_paid())
        self.listbox_preview_self.bind("<Control-Up>", lambda e: (self.move_self_paid_items(-1), "break")[1])
        self.listbox_preview_self.bind("<Control-Down>", lambda e: (self.move_self_paid_items(1), "break")[1])

        move_frame = ctk.CTkFrame(workspace_frame, width=38, fg_color="transparent", corner_radius=4)
        move_frame.pack(side="right", fill="y", padx=(8, 0))
        move_frame.pack_propagate(False)
        ctk.CTkButton(
            move_frame, text="↑", width=30, height=28, font=(UI_FONT, 16, "bold"),
            fg_color="#4A5568", hover_color="#2D3748", command=lambda: self.move_self_paid_items(-1),
        ).pack(padx=2, pady=(4, 2))
        ctk.CTkButton(
            move_frame, text="↓", width=30, height=28, font=(UI_FONT, 16, "bold"),
            fg_color="#4A5568", hover_color="#2D3748", command=lambda: self.move_self_paid_items(1),
        ).pack(padx=2, pady=2)

        btn_frame = ctk.CTkFrame(right_outer, fg_color="transparent")
        btn_frame.pack(fill="x", padx=15, pady=(5, 5))
        ctk.CTkButton(btn_frame, text="❌ 刪除選取", fg_color="#D9534F", hover_color="#C9302C", height=40, font=self.default_font, command=self.remove_from_self_paid).pack(side="left", expand=True, padx=(0, 5))
        ctk.CTkButton(btn_frame, text="🔄 全部清空", fg_color="#5BC0DE", hover_color="#31B0D5", height=40, font=self.default_font, command=self.clear_self_paid).pack(side="right", expand=True, padx=(5, 0))
        
        status_frame_self = ctk.CTkFrame(right_outer, fg_color="transparent")
        status_frame_self.pack(fill="x", padx=15, pady=(0, 10))
        self.lbl_count_self = ctk.CTkLabel(status_frame_self, text="當前已排定：0 項", font=(UI_FONT, 13, "bold"), text_color="#5BC0DE")
        self.lbl_count_self.pack(side="right")

        self.refresh_listbox(self.listbox_self, all_item_names, self.self_sort_var.get())

    def build_split_tab(self, parent):
        frame = ctk.CTkFrame(parent, corner_radius=15)
        frame.pack(fill="both", expand=True, padx=20, pady=20)

        ctk.CTkLabel(frame, text="自動執行：公式轉純值 ➔ 刪除空號 ➔ 拆分博仁與杏聯", font=(UI_FONT, 18, "bold")).pack(pady=(40, 20))

        warning_text = (
            "⚠️ 重要操作提醒：\n\n"
            "1. 請確保您選擇的檔案是已經「填妥受檢者名單」的團檢名單。\n"
            "2. 填寫完畢後，必須在 Excel 中【按下儲存】再關閉，程式才能讀取到「@」符號！\n"
            "3. 點擊執行後，程式會自動依照您原本的 VBA 巨集邏輯，產出兩份全新的 Excel 檔案。"
        )
        ctk.CTkLabel(frame, text=warning_text, font=(UI_FONT, 15), text_color="#FF6B6B", justify="left").pack(pady=20)

        file_frame = ctk.CTkFrame(frame, fg_color="transparent")
        file_frame.pack(fill="x", padx=50, pady=20)

        self.entry_split_file = ctk.CTkEntry(file_frame, font=self.default_font, height=40)
        self.entry_split_file.pack(side="left", fill="x", expand=True, padx=(0, 10))

        ctk.CTkButton(file_frame, text="📂 瀏覽檔案", font=self.default_font, height=40, fg_color="#6C757D", hover_color="#5A6268", command=self.browse_split_file).pack(side="right")

        btn_run_split = ctk.CTkButton(
            frame, text="⚡ 一鍵執行拆分巨集", font=(UI_FONT, 16, "bold"),
            fg_color="#0275D8", hover_color="#025AA5", height=50, corner_radius=10, command=self.run_split_macro
        )
        btn_run_split.pack(fill="x", padx=50, pady=(20, 0))

    def refresh_listbox(self, listbox, items, sort_by="預設 (原始順序)"):
        listbox.delete(0, tk.END)
        sorted_items = get_sorted_items_list(items, sort_by)
        for item in sorted_items:
            listbox.insert(tk.END, item)

    def filter_items(self, entry_widget, listbox_widget, sort_var):
        search_txt = entry_widget.get().strip().upper()
        sort_by = sort_var.get()
        if not search_txt:
            self.refresh_listbox(listbox_widget, all_item_names, sort_by)
            return
        filtered = [name for name in all_item_names if search_txt in name.upper()]
        self.refresh_listbox(listbox_widget, filtered, sort_by)

    def on_pub_sort_changed(self, choice):
        self.filter_items(self.entry_search_pub, self.listbox_pub, self.pub_sort_var)
        
    def on_self_sort_changed(self, choice):
        self.filter_items(self.entry_search_self, self.listbox_self, self.self_sort_var)

    @staticmethod
    def _move_picker_selection(listbox, direction):
        """讓項目選擇清單可用上下鍵移動，並維持單一明確的作用項目。"""
        size = listbox.size()
        if size <= 0:
            return "break"
        selected = listbox.curselection()
        if selected:
            current = selected[0]
        else:
            current = 0 if direction > 0 else size - 1
        target = max(0, min(size - 1, current + direction)) if selected else current
        listbox.selection_clear(0, tk.END)
        listbox.selection_set(target)
        listbox.activate(target)
        listbox.see(target)
        return "break"

    @staticmethod
    def _focus_picker_from_search(event, listbox):
        """搜尋框按下 Down 時把焦點交給第一個結果，方便一路用鍵盤操作。"""
        if listbox.size() <= 0:
            return "break"
        listbox.focus_set()
        if not listbox.curselection():
            listbox.selection_set(0)
            listbox.activate(0)
        listbox.see(listbox.index(tk.ACTIVE))
        return "break"

    def _bind_picker_keyboard(self, listbox, add_command, search_entry):
        listbox.bind("<Up>", lambda event: self._move_picker_selection(listbox, -1))
        listbox.bind("<Down>", lambda event: self._move_picker_selection(listbox, 1))
        listbox.bind("<Return>", lambda event: (add_command(), "break")[1])
        search_entry.bind("<Down>", lambda event: self._focus_picker_from_search(event, listbox))

    @staticmethod
    def _reorder_sequence(sequence, selected_indices, direction):
        """在同一清單內移動選取項目；支援多選且保留相對順序。"""
        selected_indices = sorted(set(selected_indices))
        if not selected_indices or not sequence:
            return
        selected_ids = {id(sequence[index]) for index in selected_indices}
        if direction < 0:
            for index in selected_indices:
                if index > 0 and id(sequence[index - 1]) not in selected_ids:
                    sequence[index - 1], sequence[index] = sequence[index], sequence[index - 1]
        else:
            for index in reversed(selected_indices):
                if index < len(sequence) - 1 and id(sequence[index + 1]) not in selected_ids:
                    sequence[index + 1], sequence[index] = sequence[index], sequence[index + 1]

    @staticmethod
    def _restore_listbox_selection(listbox, selected_values):
        selected_values = set(selected_values)
        listbox.selection_clear(0, tk.END)
        for index in range(listbox.size()):
            item = listbox.get(index)
            if item in selected_values:
                listbox.selection_set(index)
                listbox.activate(index)
        selected = listbox.curselection()
        if selected:
            listbox.see(selected[0])

    def move_public_items(self, direction):
        selected_indices = list(self.listbox_preview_pub.curselection())
        if not selected_indices:
            return
        selected_ids = {id(self.public_items[index]) for index in selected_indices if index < len(self.public_items)}
        self._reorder_sequence(self.public_items, selected_indices, direction)
        self.update_preview_pub()
        for index, item in enumerate(self.public_items):
            if id(item) in selected_ids:
                self.listbox_preview_pub.selection_set(index)
        selected = self.listbox_preview_pub.curselection()
        if selected:
            self.listbox_preview_pub.see(selected[0])

    def move_self_paid_items(self, direction):
        selected_indices = list(self.listbox_preview_self.curselection())
        source_items = getattr(self, "_preview_self_source_items", [])
        if not selected_indices or not source_items:
            return
        selected_sources = [source_items[index] for index in selected_indices if index < len(source_items)]
        packages = {package for package, _item_name in selected_sources}
        if len(packages) != 1:
            messagebox.showwarning("排序提示", "請先選擇同一個自費套組內的項目，再進行上下移動。")
            return

        package = next(iter(packages))
        package_items = self.self_paid_config.get(package, [])
        selected_names = {item_name for _package, item_name in selected_sources}
        selected_indices_in_package = [
            index for index, item in enumerate(package_items) if item.get("item") in selected_names
        ]
        self._reorder_sequence(package_items, selected_indices_in_package, direction)
        self.update_preview_self()
        for index, source in enumerate(getattr(self, "_preview_self_source_items", [])):
            if source in selected_sources:
                self.listbox_preview_self.selection_set(index)
        selected = self.listbox_preview_self.curselection()
        if selected:
            self.listbox_preview_self.see(selected[0])

    def add_to_public(self):
        selected_indices = self.listbox_pub.curselection()
        if not selected_indices: return
        gender_rule = self.pub_gender_var.get()
        for idx in selected_indices:
            val = self.listbox_pub.get(idx)
            if not any(x["item"] == val for x in self.public_items):
                self.public_items.append({"item": val, "gender": gender_rule})
        self.update_preview_pub()

    def remove_from_public(self):
        selected_indices = self.listbox_preview_pub.curselection()
        if not selected_indices: return
        for idx in reversed(selected_indices):
            source_items = getattr(self, "_preview_public_source_items", [])
            if 0 <= idx < len(source_items):
                item_name = source_items[idx]["item"]
                self.public_items = [x for x in self.public_items if x["item"] != item_name]
        self.update_preview_pub()

    def add_to_self_paid(self):
        pkg = self.entry_pkg.get().strip().upper()
        l_pkg = self.entry_l_pkg.get().strip().upper() # 🌟 新增抓取獨立的 L 欄代碼
        
        if not pkg:
            messagebox.showwarning("提示", "請先設定套組名稱！")
            return
        selected_indices = self.listbox_self.curselection()
        if not selected_indices: return
        gender_rule = self.self_paid_gender_var.get()
        
        # ✨ 隱性觸發：只要單項加選代碼有填寫內容，就自動啟用雙重公式 (取代打勾)
        check_l = True if l_pkg else False
        
        if pkg not in self.self_paid_config:
            self.self_paid_config[pkg] = []
        for idx in selected_indices:
            val = self.listbox_self.get(idx)
            if not any(x["item"] == val for x in self.self_paid_config[pkg]):
                # 🌟 將 l_pkg 一併存入設定
                self.self_paid_config[pkg].append({"item": val, "gender": gender_rule, "check_l_col": check_l, "l_pkg": l_pkg})
        self.update_preview_self()

    def remove_from_self_paid(self):
        selected_indices = self.listbox_preview_self.curselection()
        if not selected_indices: return
        for idx in reversed(selected_indices):
            source_items = getattr(self, "_preview_self_source_items", [])
            if 0 <= idx < len(source_items):
                pkg, item_name = source_items[idx]
                if pkg in self.self_paid_config:
                    self.self_paid_config[pkg] = [x for x in self.self_paid_config[pkg] if x["item"] != item_name]
                    if not self.self_paid_config[pkg]:
                        del self.self_paid_config[pkg]
        self.update_preview_self()

    def clear_public(self):
        self.public_items.clear()
        self.update_preview_pub()

    def clear_self_paid(self):
        self.self_paid_config.clear()
        self.update_preview_self()

    def update_preview_pub(self):
        self.listbox_preview_pub.delete(0, tk.END)
        self._preview_public_source_items = []
        for obj in self.public_items:
            tag = ""
            if obj["gender"] == "Male": tag = " 【限男】"
            elif obj["gender"] == "Female": tag = " 【限女】"
            display_item, overridden = self.get_routed_item_display_label(obj["item"])
            self.listbox_preview_pub.insert(tk.END, f"{display_item}{tag}")
            index = self.listbox_preview_pub.size() - 1
            self._preview_public_source_items.append(obj)
            self._mark_routed_listbox_item(self.listbox_preview_pub, index, overridden)
            
        if hasattr(self, 'lbl_count_pub'):
            self.lbl_count_pub.configure(text=f"當前已排定：{len(self.public_items)} 項")
        self.listbox_preview_pub.see(tk.END)

    def update_preview_self(self):
        self.listbox_preview_self.delete(0, tk.END)
        self._preview_self_source_items = []
        total_items = 0
        for pkg, items in self.self_paid_config.items():
            for obj in items:
                tag = ""
                # 🌟 視覺化顯示獨立的單項代碼 (支援隱性觸發)
                l_p = obj.get("l_pkg", "")
                if l_p:
                    tag += f" ⚡[單項:{l_p}]"
                elif obj.get("check_l_col"): 
                    tag += " ⚡[含單項]" 
                
                if obj["gender"] == "Male": tag += " 【限男】"
                elif obj["gender"] == "Female": tag += " 【限女】"
                display_item, overridden = self.get_routed_item_display_label(obj["item"])
                self.listbox_preview_self.insert(tk.END, f"[{pkg}] {display_item}{tag}")
                index = self.listbox_preview_self.size() - 1
                self._preview_self_source_items.append((pkg, obj["item"]))
                self._mark_routed_listbox_item(self.listbox_preview_self, index, overridden)
                total_items += 1
                
        if hasattr(self, 'lbl_count_self'):
            self.lbl_count_self.configure(text=f"當前已排定：{total_items} 項")
        self.listbox_preview_self.see(tk.END)

    def refresh_routing_visuals(self):
        """院所或規則改變後，同步更新兩個頁面的覆寫標籤與黃色標示。"""
        if hasattr(self, "listbox_pub"):
            self.filter_items(self.entry_search_pub, self.listbox_pub, self.pub_sort_var)
        if hasattr(self, "listbox_self"):
            self.filter_items(self.entry_search_self, self.listbox_self, self.self_sort_var)
        if hasattr(self, "listbox_preview_pub"):
            self.update_preview_pub()
        if hasattr(self, "listbox_preview_self"):
            self.update_preview_self()

    def save_config(self):
        data = {"public_items": self.public_items, "self_paid_config": self.self_paid_config}
        file_path = filedialog.asksaveasfilename(
            defaultextension=".json", filetypes=[("JSON 設定檔", "*.json")],
            initialfile="團檢預設組合.json", title="儲存當前項目設定"
        )
        if file_path:
            try:
                with open(file_path, 'w', encoding='utf-8') as f:
                    json.dump(data, f, ensure_ascii=False, indent=4)
                messagebox.showinfo("成功", "✅ 設定檔已成功儲存！")
            except Exception as e: messagebox.showerror("錯誤", "存檔失敗：" + str(e))

    def _upgrade_label(self, old_label):
        db_item = find_db_item_robust_global(old_label)
        if db_item:
            org_display = db_item['org_display']
            internal_code = db_item['internal_code']
            full_name = db_item['full_name']
            en_name = db_item['sys_code']
            category = db_item.get('category', '')
            
            if internal_code: lbl = f"[{org_display}] {internal_code} - {full_name} ({en_name})"
            else: lbl = f"[{org_display}] {full_name} ({en_name})"
                
            if category: lbl += f" | {category}" 
            return lbl
        return old_label

    def load_config(self):
        file_path = filedialog.askopenfilename(filetypes=[("JSON 設定檔", "*.json")], title="讀取歷史設定")
        if file_path:
            try:
                with open(file_path, 'r', encoding='utf-8') as f:
                    data = json.load(f)
                
                pub_data = data.get("public_items", [])
                for x in pub_data:
                    item_obj = {"item": x, "gender": "All"} if isinstance(x, str) else x
                    if isinstance(item_obj.get("gender"), bool): item_obj["gender"] = "All"
                    item_obj["item"] = self._upgrade_label(item_obj["item"])
                    
                    if not any(existing["item"] == item_obj["item"] for existing in self.public_items):
                        self.public_items.append(item_obj)
                        
                self_paid_data = data.get("self_paid_config", {})
                for pkg, items in self_paid_data.items():
                    if pkg not in self.self_paid_config:
                        self.self_paid_config[pkg] = []
                    for x in items:
                        # 🌟 相容舊版存檔，自動補上 l_pkg 空字串
                        item_obj = {"item": x, "gender": "All", "check_l_col": False, "l_pkg": ""} if isinstance(x, str) else x
                        if "l_pkg" not in item_obj: item_obj["l_pkg"] = "" 
                        if isinstance(item_obj.get("gender"), bool): item_obj["gender"] = "All"
                        item_obj["item"] = self._upgrade_label(item_obj["item"])
                        
                        if not any(existing["item"] == item_obj["item"] for existing in self.self_paid_config[pkg]):
                            self.self_paid_config[pkg].append(item_obj)
                            
                self.update_preview_pub()
                self.update_preview_self()
                if hasattr(self, "lbl_loaded_preset"):
                    self.lbl_loaded_preset.configure(
                        text=f"📄 已載入預設組合：{os.path.basename(file_path)}"
                    )
                messagebox.showinfo("成功", "✅ 設定檔載入成功！項目已自動「疊加」至當前清單中。")
            except Exception as e: messagebox.showerror("錯誤", "讀取失敗：" + str(e))

    def browse_split_file(self):
        filepath = filedialog.askopenfilename(
            filetypes=[("Excel 檔案", "*.xlsx *.xls *.xlsm")], title="選擇已填妥的團檢大表"
        )
        if filepath:
            self.entry_split_file.delete(0, tk.END)
            self.entry_split_file.insert(0, filepath)

    def run_split_macro(self):
        filepath = self.entry_split_file.get().strip()
        if not os.path.exists(filepath):
            messagebox.showwarning("錯誤", "找不到檔案，請重新確認路徑！")
            return

        try:
            base_dir = os.path.dirname(filepath)
            filename_without_ext = os.path.splitext(os.path.basename(filepath))[0]
            out_bo = os.path.join(base_dir, f"{filename_without_ext}_博仁.xlsx")
            out_xing = os.path.join(base_dir, f"{filename_without_ext}_杏聯.xlsx")

            try:
                with open(filepath, 'a'): pass
                if os.path.exists(out_bo):
                    with open(out_bo, 'a'): pass
                if os.path.exists(out_xing):
                    with open(out_xing, 'a'): pass
            except PermissionError:
                messagebox.showwarning("檔案佔用中", "⚠️ 無法執行！\n請先確認您的 Excel 是否還開著「大表」或「博仁/杏聯」的檔案。\n請將它們關閉後，再點擊執行！")
                return

            wb_src_val = openpyxl.load_workbook(filepath, data_only=True)
            ws_src_val = wb_src_val.active

            wb_src_fml = openpyxl.load_workbook(filepath, data_only=False)
            ws_src_fml = wb_src_fml.active

            # ==========================================
            # 1. 🤖 AI 動態定位：自動尋找表格偏移後的正確座標
            # ==========================================
            base_header_row = 4  # 預設為第4列 (欄位名稱列)
            name_col_idx = 8     # 預設姓名為第8欄
            dynamic_start_col = 12 # 預設檢驗項目從第12欄開始

            date_col_idx = 1
            inst_col_idx = 3
            seq_col_idx = 4

            # 往下掃描前 15 列，尋找包含「姓名」、「身分證」等關鍵字的「主標頭列」
            for r in range(1, 16):
                row_vals = [str(ws_src_val.cell(row=r, column=c).value or "").strip() for c in range(1, min(30, ws_src_val.max_column + 1))]
                if any("姓名" in v for v in row_vals) and any("身分證" in v for v in row_vals):
                    base_header_row = r
                    break

            # 精準定位各個固定欄位的位置 (就醫日期、院所、流水號、姓名...)
            for c in range(1, min(30, ws_src_val.max_column + 1)):
                val = str(ws_src_val.cell(row=base_header_row, column=c).value or "").strip()
                if "姓名" in val: name_col_idx = c
                elif "就醫日期" in val: date_col_idx = c
                elif "院所" in val: inst_col_idx = c
                elif "流水號" in val: seq_col_idx = c

            # 動態找出「自費/公費項目」是從哪一欄開始的 (掃描標頭的下一欄，通常是在前 20 欄內)
            for c in range(max(1, name_col_idx), ws_src_val.max_column + 1):
                top_tag_row = max(1, base_header_row - 3)
                top_val = str(ws_src_val.cell(row=top_tag_row, column=c).value or "").strip()
                if top_val in ["杏1", "杏2", "博1", "博2", "E", "9", "C1", "C2", "B1", "B2"]:
                    dynamic_start_col = c
                    break
            else:
                # 備用方案：如果最上層沒有標示，就找標題叫 A011 或 001 的地方
                for c in range(max(1, name_col_idx), ws_src_val.max_column + 1):
                    code_val = str(ws_src_val.cell(row=max(1, base_header_row - 1), column=c).value or "").strip()
                    if code_val and (code_val.isdigit() or code_val.startswith("A0")):
                        dynamic_start_col = c
                        break

            # 動態尋找「場次」與「院所」的全域變數 (原本在 E2, H2)
            session_code = ""
            hospital_code = ""
            for r in range(1, base_header_row):
                for c in range(1, dynamic_start_col):
                    val = str(ws_src_val.cell(row=r, column=c).value or "").strip()
                    if "場次" in val:
                        session_code = str(ws_src_val.cell(row=r, column=c+1).value or "").strip()
                        if not session_code: session_code = str(ws_src_val.cell(row=r, column=c+2).value or "").strip()
                    if "院所" in val and not hospital_code:
                        hospital_code = str(ws_src_val.cell(row=r, column=c+1).value or "").strip()
                        if not hospital_code: hospital_code = str(ws_src_val.cell(row=r, column=c+2).value or "").strip()

            # 如果還是找不到場次代碼，跳出視窗詢問
            if not session_code:
                dialog = ctk.CTkInputDialog(
                    text="系統無法自動定位「場次代碼」位置！\n\n請手動輸入本次的「場次代碼」(例如: 61)：",
                    title="⚠️ 缺少場次代碼"
                )
                user_input = dialog.get_input()

                if not user_input or not user_input.strip():
                    messagebox.showwarning("已取消", "您未輸入場次代碼，拆分作業已終止！")
                    wb_src_val.close()
                    wb_src_fml.close()
                    return

                session_code = user_input.strip()

            today = datetime.date.today()
            tw_year = today.year - 1911
            fallback_date = f"{tw_year}{today.month:02d}{today.day:02d}"

            # 樣式設定
            font_header = Font(name="Microsoft JhengHei", size=10, bold=True)
            font_normal = Font(name="Microsoft JhengHei", size=10)
            align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
            align_data = Alignment(horizontal="center", vertical="center", wrap_text=False)
            thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
            fill_r3 = PatternFill(start_color="FFFCE4D6", end_color="FFFCE4D6", fill_type="solid")
            fill_r4 = PatternFill(start_color="FFDDEBF7", end_color="FFDDEBF7", fill_type="solid")

            # ==========================================
            # 2. 處理博仁檔案
            # ==========================================
            wb_bo = openpyxl.Workbook()
            ws_bo = wb_bo.active
            ws_bo.title = "博仁"
            ws_bo.views.sheetView[0].showGridLines = True
            
            # 動態設定凍結窗格 (例如: L3)
            ws_bo.freeze_panes = f"{get_column_letter(dynamic_start_col)}3"

            current_inst_code = self.inst_var.get().split(" - ")[0].strip().upper()
            inst_data = self.inst_db.get(current_inst_code, {})

            force_to_xing = parse_routing_rules(inst_data.get("force_to_xinglian"))
            force_to_bo = parse_routing_rules(inst_data.get("force_to_boren"))
            top_tag_row = max(1, base_header_row - 3)
            forced_destinations = {}
            routing_changes = []
            routing_conflicts = []

            # 先為每一個動態欄建立唯一的路由結果，後面產生兩份檔案時共用，
            # 避免兩邊各自判斷而出現不一致。
            for c in range(dynamic_start_col, ws_src_val.max_column + 1):
                routing_values = column_routing_values(ws_src_val, c, base_header_row, top_tag_row)
                # 保留舊版可用「來源欄序號」指定的行為，但新介面主要應輸入代碼或名稱。
                routing_values["absolute_column"] = str(c)
                routing_values["relative_column"] = str(c - dynamic_start_col + 1)
                destination, xing_rule, bo_rule = resolve_forced_destination(
                    force_to_xing, force_to_bo, routing_values
                )

                if destination:
                    forced_destinations[c] = destination
                    # 依本次來源表頭推算原分檔結果；只有實際改變去向才列入告知。
                    # 未標明單一院所的欄位可能原本進入兩份檔案，E／9 則兩邊均略過。
                    original_destinations = []
                    col_tag = routing_values["org_tag"]
                    if col_tag not in ["博1", "博2", "E", "9"]:
                        original_destinations.append("xing")
                    if col_tag not in ["杏1", "杏2", "E", "9"]:
                        original_destinations.append("boren")
                    if original_destinations != [destination]:
                        destination_names = {"xing": "杏聯", "boren": "博仁"}
                        previous = "、".join(destination_names[value] for value in original_destinations) or "不分派"
                        item = routing_values["system_code"] or routing_values["english_code"] or routing_values["item_name"]
                        routing_changes.append(
                            f"{get_column_letter(c)} 欄 {item}：{previous} → {destination_names[destination]}"
                        )

                if destination == "xing" and bo_rule:
                    routing_conflicts.append(
                        f"第 {c} 欄（{routing_values['system_code'] or routing_values['english_code'] or routing_values['item_name']}）："
                        f"杏聯「{xing_rule}」優先於博仁「{bo_rule}」"
                    )
            if routing_conflicts:
                messagebox.showwarning(
                    "分派規則衝突",
                    f"院所 {current_inst_code} 的規則同時命中同一欄，已依杏聯優先處理：\n\n"
                    + "\n".join(routing_conflicts[:10])
                    + ("\n…" if len(routing_conflicts) > 10 else ""),
                )

            bo_keep_cols = list(range(1, dynamic_start_col))
            for c in range(dynamic_start_col, ws_src_val.max_column + 1):
                col_tag = str(ws_src_val.cell(row=top_tag_row, column=c).value or "").strip()
                forced_destination = forced_destinations.get(c)
                if forced_destination == "boren":
                    bo_keep_cols.append(c)
                elif forced_destination == "xing":
                    continue
                elif col_tag not in ["杏1", "杏2", "E", "9"]:
                    bo_keep_cols.append(c)

            bo_current_row = 1
            # 複製表頭 (系統代碼列、欄位名稱列)
            for src_r in [max(1, base_header_row - 1), base_header_row]:
                for new_c, src_c in enumerate(bo_keep_cols, start=1):
                    cell_src = ws_src_val.cell(row=src_r, column=src_c)
                    cell_new = ws_bo.cell(row=bo_current_row, column=new_c, value=cell_src.value)
                    cell_new.font = font_header
                    cell_new.border = thin_border
                    cell_new.alignment = align_center
                    if bo_current_row == 1: cell_new.fill = fill_r3
                    elif bo_current_row == 2: cell_new.fill = fill_r4
                bo_current_row += 1

            serial_counter_bo = 1
            for src_r in range(base_header_row + 1, ws_src_val.max_row + 1):
                name_val = str(ws_src_val.cell(row=src_r, column=name_col_idx).value or "").strip()
                if name_val == "空號" or not name_val:
                    continue

                for new_c, src_c in enumerate(bo_keep_cols, start=1):
                    val = ws_src_val.cell(row=src_r, column=src_c).value

                    if src_c == date_col_idx and (val is None or val == ""): val = fallback_date
                    if src_c == inst_col_idx and (val is None or val == ""): val = hospital_code
                    if src_c == seq_col_idx:
                        fml_cell = ws_src_fml.cell(row=src_r, column=src_c).value
                        if fml_cell is not None and str(fml_cell).strip().startswith("="):
                            val = f"{session_code}{serial_counter_bo:04d}"
                        else:
                            src_val = ws_src_val.cell(row=src_r, column=src_c).value
                            val = str(src_val).strip() if src_val is not None else ""

                    if src_c >= dynamic_start_col:
                        if val is None: val = ""

                    cell_new = ws_bo.cell(row=bo_current_row, column=new_c, value=val)
                    cell_new.font = font_normal
                    cell_new.border = thin_border
                    cell_new.alignment = align_data

                ws_bo.row_dimensions[bo_current_row].height = 20
                serial_counter_bo += 1
                bo_current_row += 1

            for col in ws_bo.columns:
                col_letter = col[0].column_letter
                if col_letter == 'K':
                    ws_bo.column_dimensions[col_letter].width = 23
                    continue
                max_length = 0
                for cell in col:
                    if cell.value:
                        val_str = str(cell.value)
                        c_len = sum(2.2 if ord(char) > 255 else 1.1 for char in val_str)
                        if c_len > max_length: max_length = c_len
                ws_bo.column_dimensions[col_letter].width = max(max_length + 1.5, 11)

            wb_bo.save(out_bo)
            wb_bo.close()

            # ==========================================
            # 3. 處理杏聯檔案
            # ==========================================
            wb_xing = openpyxl.Workbook()
            ws_xing = wb_xing.active
            ws_xing.title = "杏聯"
            ws_xing.views.sheetView[0].showGridLines = True
            
            # 動態設定凍結窗格
            ws_xing.freeze_panes = f"{get_column_letter(dynamic_start_col)}3"

            xing_keep_cols = list(range(1, dynamic_start_col))
            for c in range(dynamic_start_col, ws_src_val.max_column + 1):
                col_tag = str(ws_src_val.cell(row=top_tag_row, column=c).value or "").strip()
                forced_destination = forced_destinations.get(c)
                if forced_destination == "xing":
                    xing_keep_cols.append(c)
                elif forced_destination == "boren":
                    continue
                elif col_tag not in ["博1", "博2", "E", "9"]:
                    xing_keep_cols.append(c)

            xing_current_row = 1
            for src_r in [max(1, base_header_row - 1), base_header_row]:
                for new_c, src_c in enumerate(xing_keep_cols, start=1):
                    cell_src = ws_src_val.cell(row=src_r, column=src_c)
                    cell_new = ws_xing.cell(row=xing_current_row, column=new_c, value=cell_src.value)
                    cell_new.font = font_header
                    cell_new.border = thin_border
                    cell_new.alignment = align_center
                    if xing_current_row == 1: cell_new.fill = fill_r3
                    elif xing_current_row == 2: cell_new.fill = fill_r4
                xing_current_row += 1

            serial_counter_xing = 1
            for src_r in range(base_header_row + 1, ws_src_val.max_row + 1):
                name_val = str(ws_src_val.cell(row=src_r, column=name_col_idx).value or "").strip()
                if name_val == "空號" or not name_val:
                    continue

                for new_c, src_c in enumerate(xing_keep_cols, start=1):
                    val = ws_src_val.cell(row=src_r, column=src_c).value

                    if src_c == date_col_idx and (val is None or val == ""): val = fallback_date
                    if src_c == inst_col_idx and (val is None or val == ""): val = hospital_code
                    if src_c == seq_col_idx:
                        fml_cell = ws_src_fml.cell(row=src_r, column=src_c).value
                        if fml_cell is not None and str(fml_cell).strip().startswith("="):
                            val = f"{session_code}{serial_counter_xing:04d}"
                        else:
                            src_val = ws_src_val.cell(row=src_r, column=src_c).value
                            val = str(src_val).strip() if src_val is not None else ""

                    if src_c >= dynamic_start_col:
                        if val is None: val = ""

                    cell_new = ws_xing.cell(row=xing_current_row, column=new_c, value=val)
                    cell_new.font = font_normal
                    cell_new.border = thin_border
                    cell_new.alignment = align_data

                ws_xing.row_dimensions[xing_current_row].height = 20
                serial_counter_xing += 1
                xing_current_row += 1

            for col in ws_xing.columns:
                col_letter = col[0].column_letter
                if col_letter == 'K':
                    ws_xing.column_dimensions[col_letter].width = 23
                    continue
                max_length = 0
                for cell in col:
                    if cell.value:
                        val_str = str(cell.value)
                        c_len = sum(2.2 if ord(char) > 255 else 1.1 for char in val_str)
                        if c_len > max_length: max_length = c_len
                ws_xing.column_dimensions[col_letter].width = max(max_length + 1.5, 11)

            wb_xing.save(out_xing)
            wb_xing.close()
            wb_src_val.close()
            wb_src_fml.close()

            success_message = (
                "🎉 已成功拆分為兩個無公式的純值檔案：\n\n1. "
                + os.path.basename(out_bo) + "\n2. " + os.path.basename(out_xing)
            )
            if routing_changes:
                success_message += (
                    f"\n\n院所 {current_inst_code} 已強制變更分派 {len(routing_changes)} 項：\n"
                    + "\n".join(routing_changes[:10])
                )
                if len(routing_changes) > 10:
                    success_message += f"\n另有 {len(routing_changes) - 10} 項。"
            messagebox.showinfo("拆分成功", success_message)
        except Exception as e:
            error_details = traceback.format_exc()
            messagebox.showerror("執行錯誤", "拆分過程發生阻礙，錯誤代碼如下：\n\n" + error_details)
