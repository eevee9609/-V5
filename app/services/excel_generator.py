"""標準團檢 Excel 模板的產生流程。"""

import os
import re
from copy import copy
from datetime import datetime
from tkinter import filedialog, messagebox

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.config import APP_ROOT
from app.data import find_db_item_robust
from app.utils.dates import convert_to_taiwan_date


CURRENT_DIR = str(APP_ROOT)
find_db_item_robust_global = find_db_item_robust


_PACKAGE_CODE_SEPARATOR = re.compile(r"[,，;；\s\u3000]+")
_INVALID_FILENAME_CHARS = re.compile(r'[<>:"/\\|?*\x00-\x1f]')
_LEGACY_VERSIONED_PACKAGE_CODE = re.compile(r"^[A-Z]+\d+$")


def _gender_code_for_excel(value: object) -> object:
    """將匯入名單常見的性別文字／代碼寫成模板使用的數字 1、2。"""
    raw_text = str(value or "").strip()
    lookup_text = raw_text.casefold().replace(" ", "").replace("　", "")
    if lookup_text in {"1", "1.0", "男", "男性", "男生", "m", "male"}:
        return 1
    if lookup_text in {"2", "2.0", "女", "女性", "女生", "f", "female"}:
        return 2
    return raw_text


def _split_package_codes(raw_value: object) -> list[str]:
    """將 K／L 欄可接受的分隔符統一解析為不重複的完整代碼清單。"""
    codes: list[str] = []
    seen: set[str] = set()
    for value in _PACKAGE_CODE_SEPARATOR.split(str(raw_value or "")):
        # 允許設定端也沿用 V1.／V7. 的舊習慣；句點是 K 欄的分隔符，
        # 不屬於代碼本身，因此比對前移除結尾句點。
        code = value.strip().rstrip(".．。")
        if code and code not in seen:
            codes.append(code)
            seen.add(code)
    return codes


def _excel_exact_code_match(cell_ref: str, code: object) -> str:
    """建立 L 欄完整代碼比對公式，避免 B1 誤命中 B10。"""
    escaped_code = str(code).strip().upper().replace('"', '""')
    normalized_cell = cell_ref
    for source, replacement in (
        ('"，"', '","'),
        ('"；"', '","'),
        ('";"', '","'),
        ("CHAR(13)", '","'),
        ("CHAR(10)", '","'),
        ("CHAR(9)", '","'),
        ('"　"', '","'),
        ('" "', '","'),
    ):
        normalized_cell = f"SUBSTITUTE({normalized_cell},{source},{replacement})"
    return f'ISNUMBER(FIND(",{escaped_code},",","&UPPER({normalized_cell})&","))'


def _excel_legacy_k_code_match(cell_ref: str, code: object) -> str:
    """建立舊模板使用的 K 欄搜尋公式。

    K 欄沿用 ABCD. 這種緊湊輸入，因此 A／B／C／D／炎等代碼採局部
    搜尋。V1、V7 這類有數字的代碼則要求結尾為句點（逗號與全形句點
    會先統一），使 V1 不會誤判為 V10。
    """
    normalized_code = str(code or "").strip().upper().rstrip(".．。")
    if not normalized_code:
        return "FALSE"

    escaped_code = normalized_code.replace('"', '""')
    normalized_cell = cell_ref
    for source, replacement in (
        ('"．"', '"."'),
        ('"。"', '"."'),
        ('"，"', '"."'),
        ('","', '"."'),
        ('"；"', '"."'),
        ('";"', '"."'),
        ("CHAR(13)", '"."'),
        ("CHAR(10)", '"."'),
    ):
        normalized_cell = f"SUBSTITUTE({normalized_cell},{source},{replacement})"
    normalized_cell = f"UPPER({normalized_cell})"

    if _LEGACY_VERSIONED_PACKAGE_CODE.fullmatch(normalized_code):
        return f'ISNUMBER(FIND("{escaped_code}.",{normalized_cell}&"."))'
    return f'ISNUMBER(FIND("{escaped_code}",{normalized_cell}))'


def _excel_k_package_code_match(cell_ref: str, code: object) -> str:
    """K 欄固定使用舊模板的緊湊搜尋規則。"""
    return _excel_legacy_k_code_match(cell_ref, code)


def _safe_filename_component(value: object, fallback: str = "未指定院所") -> str:
    """將院所代號轉為 Windows 可用的檔名片段。"""
    cleaned = _INVALID_FILENAME_CHARS.sub("_", str(value or "").strip()).rstrip(". ")
    return cleaned or fallback


def _choose_output_path(inst_code: object) -> str | None:
    """詢問輸出位置，交由 Windows 另存視窗確認是否覆蓋既有檔案。"""
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    initial_filename = f"團檢模板_{_safe_filename_component(inst_code)}_{timestamp}.xlsx"
    selected_path = filedialog.asksaveasfilename(
        title="另存團檢模板",
        initialdir=CURRENT_DIR,
        initialfile=initial_filename,
        defaultextension=".xlsx",
        filetypes=[("Excel 活頁簿", "*.xlsx")],
    )
    return os.path.normpath(selected_path) if selected_path else None


def _parse_keyword_values(raw_value: object) -> list[str]:
    """將特殊公式關鍵字統一成清單，兼容舊版字串與新版陣列資料。"""
    if isinstance(raw_value, (list, tuple, set)):
        values = [str(value) for value in raw_value]
    else:
        values = str(raw_value or "").replace("，", ",").split(",")
    return [value.strip() for value in values if value.strip()]


def _get_special_formula_triggers(app_settings: object, inst_code: object) -> dict[str, list[str]]:
    """合併通用與目前院所的特殊公式規則；院所規則優先。"""
    settings = app_settings if isinstance(app_settings, dict) else {}
    merged = dict(settings.get("special_formula_keywords", {}) or {})
    by_institution = settings.get("special_formula_keywords_by_institution", {}) or {}
    if isinstance(by_institution, dict):
        # 院所代號在舊設定檔可能有大小寫或前後空白差異；產生模板時
        # 一律以標準化代號尋找，避免明明設定過院所規則卻只套用通用規則。
        normalised_code = str(inst_code or "").strip().upper()
        specific = by_institution.get(normalised_code)
        if specific is None:
            specific = next(
                (
                    value
                    for code, value in by_institution.items()
                    if str(code or "").strip().upper() == normalised_code
                ),
                {},
            )
        specific = specific or {}
        if isinstance(specific, dict):
            merged.update(specific)
    return {
        str(item): _parse_keyword_values(keywords)
        for item, keywords in merged.items()
        if str(item).strip()
    }


class ExcelGenerationMixin:

    def generate_excel(self):
        if not self.public_items and not self.self_paid_config:
            messagebox.showwarning("無法產生", "尚未排定任何公費或自費項目！")
            return
            
        try:
            try:
                reserve_val = int(self.entry_reserve.get().strip())
            except ValueError:
                reserve_val = 100 
                
            active_patients = []
            if hasattr(self, 'parsed_roster_data') and self.parsed_roster_data:
                for idx, p in enumerate(self.parsed_roster_data):
                    if self.roster_checkbox_vars[idx].get():
                        active_patients.append(p)
                        
            if len(active_patients) > 0:
                max_rows = len(active_patients)
            else:
                max_rows = reserve_val if reserve_val > 0 else 100
                
            inst_code = self.inst_var.get().split(" - ")[0].strip().upper()
            if not inst_code:
                continue_without_institution = messagebox.askyesno(
                    "尚未選擇院所",
                    "尚未選擇院所代號，產出的模板 H2 院所欄將會空白。\n\n"
                    "確定仍要繼續產生嗎？",
                    default="no",
                )
                if not continue_without_institution:
                    return
            output_path = _choose_output_path(inst_code)
            if not output_path:
                return
        
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "團檢大表"
            ws.views.sheetView[0].showGridLines = True
            ws.freeze_panes = "L5"

            font_normal = Font(name="Microsoft JhengHei", size=10) 
            font_header = Font(name="Microsoft JhengHei", size=10, bold=True)
            font_formula = Font(name="Consolas", size=10, color="FF0033CC") 
            
            fill_row1 = PatternFill(start_color="FFE2EFDA", end_color="FFE2EFDA", fill_type="solid") 
            fill_row2 = PatternFill(start_color="FFFFF2CC", end_color="FFFFF2CC", fill_type="solid") 
            fill_row3 = PatternFill(start_color="FFFCE4D6", end_color="FFFCE4D6", fill_type="solid") 
            fill_row4 = PatternFill(start_color="FFDDEBF7", end_color="FFDDEBF7", fill_type="solid") 

            # 模板標頭第 1～4 列保持單行顯示，不啟用 Excel 自動換行。
            align_center = Alignment(horizontal="center", vertical="center", wrap_text=False)
            align_data = Alignment(horizontal="center", vertical="center", wrap_text=False)
            
            thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

            row1_headers = ["A1"] * 12
            row2_headers = ["", "", "", "場次", "", "", "院所", "", "", "糞便2", "", ""]
            row3_headers = ["A011", "A001", "A003", "A002", "A004", "A006", "A008", "A005", "A007", "", "", ""]
            row4_headers = ["就醫日期", "開單日期", "院所", "流水號", "原病歷號", "身分證字號", "生日", "姓名", "性別", "飯後1", "自費", "單項BC"]

            for idx, val in enumerate(row1_headers, start=1):
                ws.cell(row=1, column=idx, value=val).font = font_header; ws.cell(row=1, column=idx).border = thin_border; ws.cell(row=1, column=idx).alignment = align_center
            for idx, val in enumerate(row2_headers, start=1):
                ws.cell(row=2, column=idx, value=val).font = font_header; ws.cell(row=2, column=idx).border = thin_border; ws.cell(row=2, column=idx).alignment = align_center
            for idx, val in enumerate(row3_headers, start=1):
                ws.cell(row=3, column=idx, value=val).font = font_header; ws.cell(row=3, column=idx).border = thin_border; ws.cell(row=3, column=idx).alignment = align_center
            for idx, name in enumerate(row4_headers, start=1):
                ws.cell(row=4, column=idx, value=name).font = font_header; ws.cell(row=4, column=idx).border = thin_border; ws.cell(row=4, column=idx).alignment = align_center

            # ⭐ [完美整合] 將輸入的院所代碼寫入 H2 儲存格 (index=8)
            if inst_code:
                ws.cell(row=2, column=8, value=inst_code).font = font_header
                ws.cell(row=2, column=8).border = thin_border
                ws.cell(row=2, column=8).alignment = align_center

            current_col_idx = 13

            for r in range(5, max_rows + 5):
                ws.row_dimensions[r].height = 20
                for c in range(1, 13):
                    cell = ws.cell(row=r, column=c)
                    cell.alignment = align_data
                    cell.border = thin_border
                    cell.font = font_normal

                formula_str = f'=IF(OR(H{r}="空號",H{r}=""),"",TEXT(TODAY(),"[$-404]emmdd"))'
                ws.cell(row=r, column=1, value=formula_str).font = font_formula
                ws.cell(row=r, column=1).number_format = '@'
                
                ws.cell(row=r, column=3, value=f'=IF(H{r}<>"", $H$2, "")').font = font_formula
                
                row_b = r - 4  
                ws.cell(row=r, column=4, value=f'=IF(OR(H{r}="空號",H{r}=""),"",$E$2&TEXT(ROW($B{row_b}),"0000"))').font = font_formula

                p_idx = r - 5
                if p_idx < len(active_patients):
                    p = active_patients[p_idx]
                    ws.cell(row=r, column=8, value=p["name"])   
                    
                    # 性別匯入可接受男／女及 1／2；模板一律寫成數字代碼。
                    ws.cell(row=r, column=9, value=_gender_code_for_excel(p.get("gender", "")))
                    
                    ws.cell(row=r, column=7, value=convert_to_taiwan_date(p["birth"]) if "birth" in p else "")
                    ws.cell(row=r, column=6, value=p["id"])  

            fixed_columns = [
                {
                    "org_display": "杏1",
                    "row2": "",
                    "internal_code": "105",
                    "sys_code": "GLU",
                    "formula": lambda r: f'=IF(AND($H{r}<>0,$K{r}<>"不抽血",$H{r}<>"空號",OR($J{r}="",$J{r}=2)),"@","")'
                },
                {
                    "org_display": "杏1",
                    "row2": "",
                    "internal_code": "106",
                    "sys_code": "PC",
                    "formula": lambda r: f'=IF(AND($H{r}<>0,$K{r}<>"不抽血",$H{r}<>"空號",ISNUMBER(FIND("1",$J{r}))),"@","")'
                }
            ]

            for col_def in fixed_columns:
                ws.cell(row=1, column=current_col_idx, value=col_def["org_display"]).fill = fill_row1
                ws.cell(row=2, column=current_col_idx, value=col_def["row2"]).fill = fill_row2
                ws.cell(row=3, column=current_col_idx, value=col_def["internal_code"]).fill = fill_row3
                ws.cell(row=4, column=current_col_idx, value=col_def["sys_code"]).fill = fill_row4
                for r in range(1, 5):
                    ws.cell(row=r, column=current_col_idx).font = font_header
                    ws.cell(row=r, column=current_col_idx).alignment = align_center
                    ws.cell(row=r, column=current_col_idx).border = thin_border
                for r in range(5, max_rows + 5):
                    c = ws.cell(row=r, column=current_col_idx, value=col_def["formula"](r))
                    c.font = font_formula
                    c.alignment = align_data
                    c.border = thin_border
                current_col_idx += 1

            # 🌟 從系統設定檔讀取通用與院所專屬特殊關鍵字規則。
            # 院所規則若與通用規則使用相同項目，則以院所設定覆蓋通用設定。
            keyword_triggers = _get_special_formula_triggers(
                getattr(self, 'app_settings', {}),
                inst_code,
            )

            missing_items = []
            routing_overrides = []

            for pkg_code, item_list in self.self_paid_config.items():
                for obj in item_list:
                    display_label = obj["item"]
                    gender_rule = obj["gender"]
                    check_l_col = obj.get("check_l_col", False) 
                    l_pkg_raw = obj.get("l_pkg", "") # 🌟 取出使用者輸入的 L 欄專用代碼
                    
                    # ✨ 隱性防呆機制：只要 l_pkg_raw 有文字，絕對啟用 L 欄公式
                    if l_pkg_raw:
                        check_l_col = True
                    
                    if inst_code in self.app_settings.get("l_col_institutions", []):
                        check_l_col = True
                    
                    db_item = find_db_item_robust_global(display_label)
                    
                    if not db_item: 
                        missing_items.append(display_label)
                        continue

                    effective_org_display, was_routed = self.get_effective_org_display(db_item)
                    if was_routed:
                        routing_overrides.append(
                            f"{db_item['internal_code'] or db_item['sys_code'] or db_item['full_name']}："
                            f"{db_item['org_display']} → {effective_org_display}"
                        )

                    full_name = db_item["full_name"]
                    is_blood_test = not any(x in full_name for x in ["(U)", "尿", "糞", "便", "抹片", "心電圖", "超音波", "分泌物"])

                    # 🌟 [智能升級] 分開處理 K 欄與 L 欄的搜尋字眼
                    pkg_aliases_k = _split_package_codes(pkg_code)
                    display_pkg_name = pkg_aliases_k[0] if pkg_aliases_k else pkg_code
                    
                    if l_pkg_raw:
                        pkg_aliases_l = _split_package_codes(l_pkg_raw)
                    else:
                        pkg_aliases_l = pkg_aliases_k # 沒填 L 代碼的話，預設跟 K 欄搜一樣的字

                    ws.cell(row=1, column=current_col_idx, value=effective_org_display).fill = fill_row1
                    ws.cell(row=2, column=current_col_idx, value=display_pkg_name).fill = fill_row2
                    ws.cell(row=3, column=current_col_idx, value=db_item["internal_code"]).fill = fill_row3
                    ws.cell(row=4, column=current_col_idx, value=db_item["sys_code"]).fill = fill_row4
                    for r in range(1, 5):
                        ws.cell(row=r, column=current_col_idx).font = font_header
                        ws.cell(row=r, column=current_col_idx).alignment = align_center
                        ws.cell(row=r, column=current_col_idx).border = thin_border

                    col_letter = get_column_letter(current_col_idx)
                    
                    extra_keywords = []
                    for key, words in keyword_triggers.items():
                        # 設定檔中的代碼一律以文字保存；資料庫舊資料偶爾可能是數字。
                        # 統一轉文字後，可讓設定頁精準選取的系統代碼確實套用到公式。
                        key_text = str(key).strip()
                        sys_code = str(db_item.get("sys_code", "") or "").strip()
                        internal_code = str(db_item.get("internal_code", "") or "").strip()
                        if key_text and (
                            key_text == sys_code
                            or key_text == internal_code
                            or key_text in full_name
                        ):
                            extra_keywords.extend(words)

                    for r in range(5, max_rows + 5):
                        find_list = []
                        
                        # 1. 專供套組 (K欄) 搜尋
                        for alias_k in pkg_aliases_k:
                            find_list.append(_excel_k_package_code_match(f"$K{r}", alias_k))
                            
                        # 2. 專供單項 (L欄) 搜尋 (若有啟用)
                        if check_l_col:
                            for alias_l in pkg_aliases_l:
                                find_list.append(_excel_exact_code_match(f"$L{r}", alias_l))

                        # 融入 UI 設定的特殊關鍵字
                        for kw in extra_keywords:
                            find_list.append(f'ISNUMBER(FIND("{kw}",$K{r}))')

                        # 將所有尋找條件用 OR 包起來
                        if len(find_list) > 1:
                            find_logic = f"OR({','.join(find_list)})"
                        else:
                            find_logic = find_list[0] if find_list else "FALSE"

                        # 組合最終條件
                        conditions = [f"$H{r}<>0", find_logic, f"$H{r}<>\"空號\""]
                        if is_blood_test: conditions.append(f"$K{r}<>\"不抽血\"")
                        if gender_rule == "Male":
                            conditions.append(f'OR($I{r}=1,$I{r}="男")')
                        elif gender_rule == "Female":
                            conditions.append(f'OR($I{r}=2,$I{r}="女")')

                        formula_str = f'=IF(AND({",".join(conditions)}),"@","")'
                        c = ws.cell(row=r, column=current_col_idx, value=formula_str)
                        c.font = font_formula
                        c.alignment = align_data
                        c.border = thin_border
                    current_col_idx += 1

            for obj in self.public_items:
                display_label = obj["item"]
                gender_rule = obj["gender"]
                
                db_item = find_db_item_robust_global(display_label)
                
                if not db_item:
                    missing_items.append(display_label)
                    continue

                effective_org_display, was_routed = self.get_effective_org_display(db_item)
                if was_routed:
                    routing_overrides.append(
                        f"{db_item['internal_code'] or db_item['sys_code'] or db_item['full_name']}："
                        f"{db_item['org_display']} → {effective_org_display}"
                    )

                full_name = db_item["full_name"]
                is_blood_test = not any(x in full_name for x in ["(U)", "尿", "糞", "便", "抹片", "心電圖", "超音波", "分泌物"])

                ws.cell(row=1, column=current_col_idx, value=effective_org_display).fill = fill_row1
                ws.cell(row=2, column=current_col_idx, value=full_name).fill = fill_row2
                ws.cell(row=3, column=current_col_idx, value=db_item["internal_code"]).fill = fill_row3
                ws.cell(row=4, column=current_col_idx, value=db_item["sys_code"]).fill = fill_row4
                for r in range(1, 5):
                    ws.cell(row=r, column=current_col_idx).font = font_header
                    ws.cell(row=r, column=current_col_idx).alignment = align_center
                    ws.cell(row=r, column=current_col_idx).border = thin_border

                for r in range(5, max_rows + 5):
                    conditions = [f"$H{r}<>0", f"$H{r}<>\"空號\""]
                    if is_blood_test: conditions.append(f"$K{r}<>\"不抽血\"")
                    if gender_rule == "Male":
                        conditions.append(f'OR($I{r}=1,$I{r}="男")')
                    elif gender_rule == "Female":
                        conditions.append(f'OR($I{r}=2,$I{r}="女")')

                    formula_str = f'=IF(AND({",".join(conditions)}),"@","")'
                    c = ws.cell(row=r, column=current_col_idx, value=formula_str)
                    c.font = font_formula
                    c.alignment = align_data
                    c.border = thin_border
                current_col_idx += 1

            for col in ws.columns:
                col_letter = col[0].column_letter
                if col_letter == 'K':
                    ws.column_dimensions[col_letter].width = 23 
                    continue
                max_length = 0
                for cell in col:
                    if cell.value and not str(cell.value).startswith("="):
                        val_str = str(cell.value)
                        c_len = sum(2.2 if ord(char) > 255 else 1.1 for char in val_str)
                        if c_len > max_length: max_length = c_len
                ws.column_dimensions[col_letter].width = max(max_length + 1.5, 11)

            wb.save(output_path)
            
            if missing_items:
                msg = f"⚠️ 警告：發現有 {len(missing_items)} 個項目已被跳過！\n因為它們在總表中不存在，且連系統代碼都對不起來。\n\n被跳過的項目：\n"
                for m in missing_items[:10]:
                    msg += f"- {m}\n"
                if len(missing_items) > 10:
                    msg += "...\n(以下省略)"
                msg += f"\n\n模板仍已產生，請在使用前確認缺少的項目：\n{output_path}"
                messagebox.showwarning("部分項目遺失", msg)
            else:
                total_self_paid = sum(len(items) for items in self.self_paid_config.values())
                summary_msg = (
                    f"🎉 完美連動！團檢模板已成功生成！\n\n"
                    f"📊 本次大表包含：\n"
                    f"- 👥 受檢者公式預劃：{max_rows} 行\n"
                    f"- 🏥 公費項目：{len(self.public_items)} 項\n"
                    f"- 💰 自費項目：{total_self_paid} 項\n\n"
                    f"- 🔄 院所強制分派覆寫：{len(routing_overrides)} 項\n\n"
                    f"儲存檔案：\n{output_path}"
                )
                messagebox.showinfo("生成成功", summary_msg)
            
        except Exception as e:
            messagebox.showerror("匯出失敗", "寫入 Excel 時發生阻礙：\n" + str(e))
